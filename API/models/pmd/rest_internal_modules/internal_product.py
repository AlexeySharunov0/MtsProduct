import os
import warnings
from datetime import datetime
from typing import Any

import allure
import openpyxl
from PIL import Image

from API.data.pmd import LifeCycle, StoryForUI, Owned, Health
from API.models.pmd.admin_role_model import RoleModelApi
from API.models.pmd.rest_internal_modules.curators import Curators
from API.models.ppinfo.brief import Brief
from API.models.ppinfo.product_v1 import ProductV1
from API.models.ppinfo.rest import Rest

warnings.filterwarnings("ignore", message="numpy.dtype size changed")

HOST = 'https://stage.prodboard.example.com'


class InternalProductUrls:
    """
    https://stage.prodboard.example.com/pmd/swagger/rest/internal/index.html#/
    """
    #### INTERNAL PRODUCT

    # Выгрузка Excel кластера
    CLUSTER_XLSX = 'v1/product/cluster'

    # Удаление документа/ссылки/лого по BI Code продукта
    DELETE_FILE_BY_ID = 'v1/product/file/{}'

    # Получение атомарных продуктов по BI Code продукта
    GET_ATOMIC = 'v1/product/{}/atomic'

    # Создание документа/ссылки/лого по BI Code продукта
    POST_FILE = 'v1/product/{}/file'

    # Получения икселя историчности по BI Code продукта
    STORY_EXCEL = 'v1/product/{}/story/excel'

    # Историчность на странице продукта
    STORY_FOR_UI = 'v1/product/{}/story'

    # Фильтры на странице продукота
    FILTERS_FOR_STORY_UI = 'v1/specification/form/attributeStory'

    # Получение жизненного цикла продукта по BI Code
    LIFE_CYCLE = 'v2/product/{}/lifeCycle'

    # Получения Моих продуктов и списко продуктов возможных для делегирования
    OWNED_PRODUCT = 'v2/product/owned'

    # Получения Здоровья продукта
    HEALTH = 'v1/product/{}/health'


class InternalProduct(InternalProductUrls, Rest, ProductV1):

    #### INTERNAL PRODUCT
    @allure.step('Получение Excel кластера')
    def get_excel_cluster(self, client) -> bytes:
        response = client.get(self.CLUSTER_XLSX, raw_format=True)

        return response.content

    @allure.step('Добавление документа')
    def add_doc(self, client, data, file_format: str = '.xlsx', product_id: str = 'BI_2687') -> str:
        file_tmp = open('filename_test' + file_format, 'w+')
        file_tmp.write('test123')
        file_tmp.close()

        with open('filename_test' + file_format, 'rb') as f:
            file_dict = {'file': f}

            client.post(self.POST_FILE.format(product_id), files=file_dict, json=data)

        os.remove('filename_test' + file_format)

        with allure.step('Проверяем что документ добавлен'):
            get_docs = self.get_docs_by_bi_code(client, product_id)

            assert any(
                'filename_test' in d['name'] for d in get_docs['documents']), f'Документы такие {get_docs["documents"]}'

            doc_id = get_docs['documents'][0]['id']

            return doc_id

    @allure.step('Добавление ссылки')
    def add_link(self, client, data: dict[str, Any], product_id: str = 'BI_2687'):
        client.post(self.POST_FILE.format(product_id), json=data, files=dict(foo='bar'))

        with allure.step('Проверяем что ссылка добавлена'):
            get_docs = self.get_docs_by_bi_code(client, product_id)

            assert any(
                'linkname_test' in d['name'] for d in get_docs['documents']), f'Документы такие {get_docs["documents"]}'

            doc_id = get_docs['documents'][0]['id']

            return doc_id

    @allure.step('Добавление логотипа')
    def add_logo(self, client, data: dict[str, Any], image_format: str = 'PNG', product_id: str = 'BI_2687') -> None:
        image = Image.new('RGBA', (500, 500), (255, 0, 0, 100))
        image.save('example.png', f'{image_format}')

        with open('example.png', 'rb') as f:
            files = {'file': ('example.png', f, f'image/{image_format}')}
            response = client.post(self.POST_FILE.format(product_id), files=files, json=data)
        os.remove('example.png')

    @allure.step('Удаление документа/ссылки/лого')
    def delete_doc(self, client, document_id: str, product_id: str = 'BI_2687') -> None:
        client.delete(self.DELETE_FILE_BY_ID.format(document_id))

        with allure.step('Проверяем что документ удален'):
            doc_after_delete = self.get_docs_by_bi_code(client, product_id)
            # doc_after_delete = client.get(self.GET_DOCUMENTS.format(product_id), **self.params,
            #                                 params={'id': product_id})

            assert document_id not in [doc['id'] for doc in
                                       doc_after_delete['documents']], f'{doc_after_delete["documents"]}'

    @allure.step('Получение атомарных продуктов')
    def get_atomic_by_code(self, client, product_id: str) -> dict:
        response = client.get(self.GET_ATOMIC.format(product_id))

        return response

    @property
    def get_life_cycle_fields(self):
        return ['(PPinfo_status)Idea_formation', '(PPinfo_status)In_progress', '(PPinfo_status)Pilot',
                '(PPinfo_status)In_use', 'endFunding', '(PPinfo_status)Out_of_use']

    @allure.step('Получение жизненного цикла по продуктам')
    def get_life_cycle(self, client, product_code: str = 'BI_2687') -> dict[str, Any]:
        response = client.get(self.LIFE_CYCLE.format(product_code))

        LifeCycle(**response)

        return response

    @allure.step('Проверяем сортировку статусов в прогрессбаре')
    def check_sorted_subset(self, main_list: list, sub_list: list) -> bool:
        it = iter(main_list)
        return all(any(x == e for x in it) for e in sub_list)

    @allure.step('Получение Моих продуктов')
    def get_owned_products(self, client) -> dict:

        response = client.get(self.OWNED_PRODUCT)

        Owned(**response)

        return response

    @allure.step('Получение Здоровья продукта')
    def get_product_health(self, client, product_code: str) -> dict:
        response = client.get(self.HEALTH.format(product_code))

        Health(health=response)
        return response



# STORY
class Story(InternalProductUrls, ProductV1, Brief, RoleModelApi, Curators):

    @property
    def get_moves(self):
        moves = {
            'string': ['string_add', 'string_replace', 'string_remove'],
            'many_select': ['manyselect_add', 'manyselect_replace', 'manyselect_remove'],
            'date': ['date_add', 'date_replace', 'date_remove'],
            'num': ['num_add', 'num_replace', 'num_remove'],
            'user': ['user_add', 'user_replace', 'user_remove'],
            'multi': ['multi_add', 'multi_replace', 'multi_remove'],
            'multistring': ['multistring_add', 'multistring_replace', 'multistring_remove'],
            'link': ['link_add', 'link_replace', 'link_remove'],
            'main_product': ['main_product_add', 'main_product_replace', 'main_product_remove']

        }

        return moves

    @allure.step('Получение историчности для UI')
    def get_story_for_ui_by_product(self, client, bi_code: str, attributes_list: list[str] | None = None,
                                    is_date_from: bool = True) -> list:
        """
        Получение историчности для отображения на фронте по продукту. В методе присутствуют возможности добавления как второй даты, так и списка атрибутов
        Args:
            bi_code: Код продукта
            attributes_list: Список атрибутов которые необходимо передать в параметры запроса  (опциональный параметр)
            is_date_from: Вторая дата (опциональный параметр)
            client: Клиент для запросов
        Returns:
            Список значений историчности для отображения на фронте
        """
        today_date = datetime.today().strftime('%d.%m.%Y')
        params = [
            ('dateTo', today_date)
        ]
        if attributes_list:
            params += [('attributeCodes', attr) for attr in attributes_list]
        if is_date_from:
            params += [('dateFrom', '01.05.2024')]

        response = client.get(self.STORY_FOR_UI.format(bi_code), params=params)

        StoryForUI(story=response)

        return response

    @allure.step('Получение фильтров историчности для UI')
    def get_filters_story_for_ui_by_product(self, client, bi_code: str) -> list:
        """
        Получение фильтров историчности для отображения на фронте по продукту.
        Args:
            bi_code: Код продукта
            client: Клиент для запросов
        Returns:
            Список фильтров историчности для отображения на фронте
        """
        params = {
            'productCode': bi_code
        }
        response = client.get(self.FILTERS_FOR_STORY_UI, params=params)

        # StoryForUI(story=response)

        return response

    @allure.step('Получения икселя историчности')
    def get_story_excel_by_product(self, client, bi_code: str, attributes_list: list[str] | None = None,
                                   is_date_from: bool = False) -> openpyxl:
        """
        Получение икселя историчности по продукту. В методе присутствуют возможности добавления как второй даты, так и списка атрибутов
        Args:
            bi_code: Код продукта
            attributes_list: Список атрибутов которые необходимо передать в параметры запроса  (опциональный параметр)
            is_date_from: Вторая дата (опциональный параметр)
        Returns:
            Подготовленный иксель историчности в открытый openpyxl
        """
        today_date = datetime.today().strftime('%d.%m.%Y')
        params = [
            ('dateTo', today_date)
        ]
        if attributes_list:
            params += [('attributeCodes', attr) for attr in attributes_list]
        if is_date_from:
            params += [('dateFrom', '01.05.2024')]

        response = client.get(self.STORY_EXCEL.format(bi_code), params=params, raw_format=True)

        with open('filename_excel.xlsx', 'wb') as f:
            f.write(response.content)

        xls = 'filename_excel.xlsx'

        wb = openpyxl.load_workbook(xls)

        return wb.active

    @allure.step('Создание продукт и изменение свойств в нем для того чтоб в историчности появились записи')
    def create_and_replace_product_admin(self, client, client_create_product, moves: list[str], for_ui: bool = False,
                                         external: bool = False):
        """
        Создание продукта админом и админом же свойств для того чтоб в историчности появились запись в историчности excel
        Args:
            moves: Действие из методов attr_storage или attr_storage_brief, то есть ключ из словаря, по которому мы берем нужный нам объект (тело запроса) для изменения свойств
            for_ui: bool переключатель для вывода историчности в Excel и для фронта. По дефолту false
            external: bool переключатель для того чтоб получить id продукта, который мы создали, для дальнейшего использования в других методах, это нужно для external api
        Returns:
            Подготовленный иксель историчности, уже, подразумевается что он будет с данными которые мы только что изменили
        """
        data = self.get_data_for_product_creation('Vnutrennii')
        product_id = self.create_product(client_create_product, data)

        for move in moves:
            data = self.get_attribute_update_body(move, bi_code=product_id)

            self.change_product(client_create_product, data)

        if external:
            return product_id

        if for_ui:
            return self.get_story_for_ui_by_product(client, product_id)
        else:
            return self.get_story_excel_by_product(client, product_id)

    @allure.step('Получение инициатора по номеру строки')
    def get_initiator_from_excel(self, excel, row: int) -> tuple[str, str]:
        """
        Мы знаем в каких столбцах отображается инициатор и его логин. В данным метод мы передаем лишь номер
        Args:
            excel: Иксель документов в котором будет осуществяться поиск
            row: Строка в которой будет осуществяться поиск
        Returns:
            Кортеж из ФИО и логина инициатора
        """
        return excel.cell(row=row, column=6).value, excel.cell(row=row, column=7).value

    @allure.step('Получение модератора по номеру строки')
    def get_moderator_from_excel(self, excel, row: int) -> tuple[str, str]:
        """
        Мы знаем в каких столбцах отображается модератор и его логин. В данным метод мы передаем лишь номер
        Args:
            excel: Иксель документов в котором будет осуществяться поиск
            row: Строка, в которой будет осуществяться поиск
        Returns:
            Кортеж из ФИО и логина модератора
        """
        return excel.cell(row=row, column=8).value, excel.cell(row=row, column=9).value

    @allure.step(
        'Создание продукт и изменение свойств в нем через заявку для того чтоб в историчности появились записи')
    def create_and_replace_product_user(self, client, brief_client_admin, brief_client_user, moves: list[str],
                                        auto_moderation: bool = False,
                                        for_ui: bool = False,
                                        external: bool = False):
        """
        Создание продукта админом и редактирование свойств юзером для того чтоб в историчности появились запись в историчности excel
        Args:
            moves: Действие из методов attr_storage или attr_storage_brief, то есть ключ из словаря, по которому мы берем нужный нам объект (тело запроса) для изменения свойств
            auto_moderation: Флаг использования автомодерации, если он включен то изменения проходят с автомодерацией
            for_ui: переключатель excel и ui
            external: bool переключатель для того чтоб получить id продукта, который мы создали, для дальнейшего использования в других методах, это нужно для external api
            client: API client
        Returns:
            Подготовленный иксель историчности, уже, подразумевается что он будет с данными которые мы только что изменили
        """
        data = self.get_data_for_product_creation('Vnutrennii')
        product_id = self.create_product(brief_client_admin, data)
        for move in moves:
            body = self.get_briefs_body(move, bi_code=product_id)
            if auto_moderation:
                self.crete_brief_by_admin(brief_client_admin, body)
            else:
                self.full_cycle(brief_client_admin, brief_client_user, body, 'statusApproval')

        if external:
            return product_id

        if for_ui:
            return self.get_story_for_ui_by_product(client, product_id)
        else:
            return self.get_story_excel_by_product(client, product_id)

    @allure.step('Изменения видимости атрибута для пользователей в ролевке')
    def change_visibility_attribute(self, client, is_visible: bool, role_id: str = '733882ff-3b13-46eb-ad4a-31d1942e7def', attribute_id: str = '629b9642-4d00-4d8a-a6f9-c570ffb1ed15'):

        data = self.get_data_for_manage(
            {'attributeId': attribute_id, 'isEditable': False, 'isView': is_visible,
             'isOwnedView': False})  # Атрибут Дата ПК

        self.manage_attr(client, role_id, data)  # Группа Пользователи в ролевке

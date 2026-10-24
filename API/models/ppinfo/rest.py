import os
import warnings
from tempfile import NamedTemporaryFile
from typing import Dict, Any, List

import allure
import openpyxl
import pandas as pd
from openpyxl.workbook import Workbook

from API.data.ppinfo import Metric, \
    AttrList, Report, Modules, EcoModule, BriefDoc
from API.models.cache.cache import Cache
from API.models.ppinfo.product_v2 import ProductV2
from API.models.ppinfo.tag import Tag
from API.utils.helpers import get_atomic_products

warnings.filterwarnings("ignore", message="numpy.dtype size changed")

HOST = 'https://stage.prodboard.example.com'


class RestUrls:
    """
    https://stage.prodboard.example.com/ppinfo/swagger/rest/index.html
    """
    #### BRIEF
    CREATE_BRIEF_DOC = 'v1/brief'

    GET_DOC_BI_ID = 'v1/brief/{}/file'

    GET_BRIEF_EXCEL = 'v1/brief/excel'

    #### PRODUCT
    METRICS = 'v2/product/{}/metrics'

    DOCUMENTS_XLSX = 'v1/product/document/xlsx'

    GET_FILE_BY_ID = 'v1/product/file/{}'

    ## pmd/internal/rest/api
    GET_DOCUMENTS = 'v1/product/{}/documents'

    # Иксель Все продукты
    GET_XLSX_ALL = 'v1/product/list/xlsx/general'

    # Иксель Все продукты, включая атомарные
    GET_XLSX_ALL_AND_ATOMIC = 'v1/product/list/xlsx/generalAndAtomics'

    # Иксель Продукты и хабы
    GET_XLSX_PRODUCTS_AND_HUBS = 'v1/product/list/xlsx/productsAndHubs'

    # Иксель Архив
    GET_XLSX_ARCHIVE = 'v1/product/list/xlsx/archive'

    GET_XLSX_UPDATE = 'v1/product/xlsx/updates'

    GET_STORY = 'v2/product/story'

    GET_ATTRIBUTES = 'v2/product/story/attributes'

    EXCEL_EPM = 'v1/product/xlsx/epm'

    ECO_MODULE = 'v1/product/{}/epm'

    EPM = 'v1/product/epm/{}'

    #### USER
    USER_REPORT = 'v1/user/report'


class Rest(RestUrls, Cache, Tag, ProductV2):

    ####################################################################  Добавление документа

    @allure.step("Получение даты для заявки с документом")
    def get_data_for_doc_brief(self, access_type: str = 'public', name: str = "test_brief") -> Dict[str, Any]:
        if not name:
            return {
                'data': '{"name":"","description":"","accessType":"public","code":"BI_3266","productId":"BI_3266","type":"document"}',
                'type': 'addDocument'
            }
        if access_type == "public":
            return {
                'data': '{"name":"test_brief","description":"","accessType":"public","code":"BI_3266","productId":"BI_3266","type":"document"}',
                'type': 'addDocument'
            }
        else:
            return {
                'data': '{"name":"test_brief","description":"","accessType":"chanelDepended","code":"BI_3266","productId":"BI_3266","type":"document"}',
                'type': 'addDocument'
            }

    @allure.step('Отправляем бриф на добавление документа админом/автосогласование')
    def post_brief_doc(self, client, body: Dict[str, Any], file_format: str = ".xlsx") -> str:
        with NamedTemporaryFile(delete=False, suffix=file_format) as file_tmp:
            file_tmp.write(b"test123")
            file_path = file_tmp.name

        try:
            with open(file_path, 'rb') as file_obj:
                file_dict = {"file": file_obj}
                response = client.post(self.CREATE_BRIEF_DOC, json=body, files=file_dict)

            BriefDoc(**response)
            return response['id']

        finally:
            os.remove(file_path)

    @allure.step('Отправляем бриф на добавление документа юзером')
    def post_brief_doc_user(self, client, body: Dict[str, Any], file_format: str = ".xlsx") -> str:
        with NamedTemporaryFile(delete=False, suffix=file_format) as file_tmp:
            file_tmp.write(b"test123")
            file_path = file_tmp.name

        try:
            with open(file_path, 'rb') as file_obj:
                file_dict = {"file": file_obj}
                response = client.post(self.CREATE_BRIEF_DOC, json=body, files=file_dict)

            BriefDoc(**response)
            return response['id']

        finally:
            os.remove(file_path)

    @allure.step('Отправляем инвалидный бриф на добавление документа')
    def invalid_post_brief_doc(self, client, body: Dict[str, Any], file_format: str = ".xlsx") -> None:
        with NamedTemporaryFile(delete=False, suffix=file_format) as file_tmp:
            file_tmp.write(b"test123")
            file_path = file_tmp.name

        try:
            with open(file_path, 'rb') as file_obj:
                file_dict = {"file": file_obj}
                client.post(self.CREATE_BRIEF_DOC, json=body, files=file_dict, expected_status_code=400)
        finally:
            os.remove(file_path)

    @allure.step('Скачиваем документ из заявки по id')
    def download_doc_from_brief(self, client, doc_id: str) -> None:

        client.get(self.GET_DOC_BI_ID.format(doc_id), raw_format=True)

    @allure.step('Получение Excel по заявкам')
    def get_brief_excel(self, client, year: str = '2025'):
        params = {
            'date': year
        }
        response = client.get(self.GET_BRIEF_EXCEL, params=params, raw_format=True)

        return response.content

    ####################################################################  Метрики

    @allure.step("Получение данных по метрикам")
    def get_metrics(self, client, code: str) -> Dict[str, Any]:
        params = {
            "id": code
        }
        response = client.get(self.METRICS.format(code), params=params)

        Metric(**response)

        return response

    #################################################################### Иксель документы

    @allure.step('Получение Excel Все документы')
    def get_all_docs_excel(self, client) -> bytes:
        response = client.get(self.DOCUMENTS_XLSX, raw_format=True)

        return response.content

    @allure.step("Создание Excel")
    def create_excel_doc(self, excel_content: bytes) -> pd.DataFrame:
        with open("filename_excel.xlsx", 'wb') as f:
            f.write(excel_content)

        return pd.read_excel("filename_excel.xlsx")

    def get_only_active_atomics(self, client_ppinfo):

        active = []
        for product_id in get_atomic_products():
            product = self.get_info_by_product_code(client_ppinfo, product_id)
            if product['status']['code'] != 'Arkhiv' and product['reasonDeletion'] == None and not product['isHub']:
                active.append(product_id)

        return active

    def get_only_active_hub_atomics(self, client_ppinfo):

        active = []
        for product_id in get_atomic_products():
            product = self.get_info_by_product_code(client_ppinfo, product_id)
            if product['status']['code'] != 'Arkhiv' and product['reasonDeletion'] == None and product['isHub']:
                active.append(product_id)

        return active

    def get_only_deleted_atomics(self, client_ppinfo):

        deleted = []
        for product_id in get_atomic_products():
            product = self.get_info_by_product_code(client_ppinfo, product_id)
            if product['reasonDeletion'] != None:
                deleted.append(product_id)

        return deleted

    ##################################################### Получение данных по документам в продукте по коду

    @allure.step("Получение документа по айди")
    def get_doc_by_id(self, client, file_id: str) -> None:

        client.get(self.GET_FILE_BY_ID.format(file_id), raw_format=True)

    #################################################################### Добавление и удаление документов

    @allure.step('Получение документов по коду BI')
    def get_docs_by_bi_code(self, client, bi_code) -> Dict[str, Any]:

        data = {
            'id': bi_code
        }

        response = client.get(self.GET_DOCUMENTS.format(bi_code), params=data)

        return response

    ##################################################### Excel продуктов

    ########################  Все продукты

    @allure.step('Получение Excel Все продукты')
    def get_excel_all(self, client) -> bytes:

        response = client.get(self.GET_XLSX_ALL, raw_format=True)

        return response.content

    @allure.step('Получение Excel Все продукты, включая атомарные')
    def get_excel_all_and_atomic(self, client) -> bytes:

        response = client.get(self.GET_XLSX_ALL_AND_ATOMIC, raw_format=True)

        return response.content

    @allure.step('Получение Excel Продукты и хабы')
    def get_excel_products_and_hubs(self, client) -> bytes:

        response = client.get(self.GET_XLSX_PRODUCTS_AND_HUBS, raw_format=True)

        return response.content

    @allure.step('Получение Excel Архив')
    def get_excel_archive(self, client) -> bytes:

        response = client.get(self.GET_XLSX_ARCHIVE, raw_format=True)

        return response.content

    @allure.step("Добавление столбца в Excel Все продукты")
    def add_col_to_excel(self, client, attr_id: str, tag_id: str = "includeIntoXlsxReport") -> None:
        try:
            self.add_attr(client, attr_id, tag_id)
        except AssertionError:
            pass

    @allure.step("Проверка соответствия столбцов с настройками")
    def check_columns(self, ppinfo_client, sheet_name: str, delete_main_product: bool = False):

        active_attributes = []
        all_attributes = self.get_tag_info(ppinfo_client, 'includeIntoXlsxReport')

        for attribute in all_attributes['attributes']:
            if attribute['isEnable']:
                active_attributes.append(attribute['name'])

        sheet = openpyxl.load_workbook('filename_excel.xlsx')[sheet_name]

        column_names = []
        if delete_main_product:
            active_attributes.remove('Основной продукт')

        for column in sheet.iter_cols(1, sheet.max_column):
            if 'e-mail' not in column[0].value.lower():
                column_names.append(column[0].value)

        set_columns = set(column_names)
        set_active = set(active_attributes)

        extra_columns = set_columns - set_active
        missing_columns = set_active - set_columns

        extra_columns_list = "\n".join([f"- {col}" for col in extra_columns])
        missing_columns_list = "\n".join([f"- {attr}" for attr in missing_columns])

        assert column_names == active_attributes, (
            "Столбцы не соответствуют настройкам!\n"
            "=== Столбцы, которые есть в Excel, но отсутствуют в активных атрибутах ===\n"
            f"{extra_columns_list}\n"
            "=== Атрибуты, которые отсутствуют в Excel, но есть в активных настройках ===\n"
            f"{missing_columns_list}"
        )
    @allure.step('Получение индекса столбца по его имени')
    def get_column_index(self, wb: Workbook, column_name: str, sheet_name: str = 'Все направления') -> int:
        """
        Получение индекса столбца по его имени
        Args:
            wb: Иксель документ
            column_name: Название столбца

        Returns:

        """
        ws = wb[sheet_name]

        headers = [cell.value for cell in ws[1]]
        if headers.count(column_name) > 1:
            column_index = headers.index(column_name) + 1
        else:
            column_index = headers.index(column_name)

        return column_index

    @allure.step('Удаление строк из икселя')
    def delete_rows_from_excel(self, wb: Workbook, column_index: int, sheet_name: str = 'Все направления', is_hub: bool = False) -> None:
        """
        Удаление лишних строк из икселя
        Args:
            wb:
            column_index:
            sheet_name:
            is_hub:
        """
        filtered_data = []

        ws = wb[sheet_name]

        for row in ws.iter_rows(min_row=1, max_row=1):
            filtered_data.append([cell.value for cell in row])


        for row in ws.iter_rows(min_row=2):
            fifth_column_value = row[column_index].value
            if is_hub:
                hub_column_index = self.get_column_index(wb, 'Тип объекта учёта', sheet_name)
                if row[hub_column_index].value == 'Хаб':
                    filtered_data.append([cell.value for cell in row])
                    continue

            if fifth_column_value is not None and fifth_column_value != 'Не определено':
                filtered_data.append([cell.value for cell in row])
        ws.delete_rows(1, ws.max_row)

        for row_data in filtered_data:
            ws.append(row_data)

        wb.save('filename.xlsx')



    @allure.step("Получение продуктов по направлениям")
    def get_products_by_directions(self, client, direction, with_atomic: str = 'false') -> list:
        """
        Получение продуктов по направлениям
        Args:
            client: Апи клиент
            direction: Направление (если нужны продукты по всем направлениям, то пишем Все направления)
            with_atomic: Если нужно чтоб в список попадали так же и атомарные продукты, то нужно указать with_atomic='true', СТРОКОЙ

        Returns:
            Список продуктов
        """
        all_derections = self.get_products_on_main_page(client, self.get_sections_codes[direction], with_atomic)

        products = []
        direction = 'Ad-tech' if direction == 'Ad-teh' else direction # костыль потому что код значения ad-tec, называется он ad-tech

        for derec in all_derections['sections']:
            if derec['name'] == direction:
                for sub_section in derec['subSections']:
                    for product in sub_section['products']:
                        products.append(product['code'])

        return products

    @allure.step("Получение хабов по направлениям")
    def get_hubs_by_directions(self, client, direction, with_atomic: str = 'false') -> list:

        all_derections = self.get_hubs_on_main_page(client, self.get_sections_codes[direction], with_atomic)

        products = []

        for derec in all_derections['sections']:
            if derec['name'] == 'Ad-tech' if direction == 'Ad-teh' else direction:
                for product in derec['products']:
                    products.append(product['code'])

        return products

    @allure.step("Получение Архива по направлениям")
    def get_archive_by_directions(self, client, direction) -> list:

        all_derections = self.get_archive_on_main_page(client)

        products = []

        for derec in all_derections:
            if derec['name'] == direction:
                for product in derec['products']:
                    products.append(product['code'])

        return products


    ##################################################### Получение Excel Изменение атрибутов по датам
    @allure.step('Получение Excel Изменение атрибутов по датам')
    def get_excel_update(self, client, params: Dict[str, Any]) -> bytes:

        response = client.get(self.GET_XLSX_UPDATE, params=params, raw_format=True)

        return response.content

    ##################################################### Получение Excel для внешних и внутренних на дату/период
    @allure.step('Получение Excel для внешних и внутренних на дату/период')
    def get_story_excel(self, client, params) -> bytes:
        response = client.get(self.GET_STORY, params=params, raw_format=True)

        return response.content

    @allure.step("Получаем дату для Excel Внутренние и Внешние на период/дату")
    def get_story_data(self, attributes: str = "", date_from: str = "", date_to: str = "") -> Dict[str, Any]:
        data = {
            'productType': 'Vneshnii',
            'from': f'{date_from}T21:00:00.000Z'
        }
        if attributes:
            data.update({"attributes": attributes})
        if date_to:
            data.update({"to": f'{date_to}T20:59:59.999Z'})
        return data

    ##################################################### Получение атрибутов

    @allure.step('Получение списка атрибутов для скачивания эксель')
    def get_attributes_list(self, client) -> List[Dict[str, str]]:
        response = client.get(self.GET_ATTRIBUTES)

        AttrList(attributes=response)

        return response

    ##################################################### Получение инфы о юзере

    @allure.step('Получение инфы о доп, уволенных и неактивных юзерах')
    def get_users_report(self, client) -> List[Dict[str, Any]]:
        response = client.get(self.USER_REPORT)

        Report(report=response)

        return response

    ##################################################### EPM / ECO-MODULES
    @property
    def get_statuses_list(self):
        return ["(PPinfo_status)In_use", "Zapushchen_v_PE", "(PPinfo_status)Pilot", "(PPinfo_status)Waiting_decision",
                "(PPinfo_status)In_progress", "(PPinfo_status)Idea_formation", "Arkhiv"]

    @allure.step("Получение экомодулей по коду BI")
    def eco_modules_by_code(self, client, bi_code: str) -> Dict[str, Any]:
        response = client.get(self.ECO_MODULE.format(bi_code))

        Modules(**response)

        return response

    @allure.step("Получение продуктов с ЭМП по коду BI")
    def epm_by_code(self, client, bi_code) -> Dict[str, Any]:
        response = client.get(self.EPM.format(bi_code))

        EcoModule(**response)

        return response

    ##################################################### Excel epm

    @allure.step("Получение excel по ЕПМ")
    def excel_epm(self, client):
        with allure.step('Получем иксель ЕПМ'):
            response = client.get(self.EXCEL_EPM, raw_format=True)
        with open('filename_excel.xlsx', 'wb') as f:
            f.write(response.content)

        return openpyxl.load_workbook('filename_excel.xlsx')
